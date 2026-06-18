import os
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.conf import settings


# ── COLORS ───────────────────────────────────────────
RED        = 'C00000'
DARK_BLUE  = '1F3864'
LIGHT_BLUE = 'D9E1F2'
YELLOW     = 'FFD966'
GREEN      = 'E2EFDA'
ORANGE     = 'FCE4D6'
GRAY       = 'F2F2F2'
WHITE      = 'FFFFFF'
TEAL       = '008080'


# ── STYLE HELPERS ────────────────────────────────────
def hfill(color):
    return PatternFill('solid', fgColor=color)

def hfont(bold=True, size=10, color='000000'):
    return Font(bold=bold, size=size,
                color=color, name='Calibri')

def halign(h='center', v='center', wrap=False):
    return Alignment(horizontal=h, vertical=v,
                     wrap_text=wrap)

def hborder():
    s = Side(style='thin', color='BFBFBF')
    return Border(left=s, right=s, top=s, bottom=s)

def sc(ws, r, c, value=None, formula=None,
       fill=None, font=None, align=None,
       border=None, fmt=None):
    cell = ws.cell(row=r, column=c)
    if formula:
        cell.value = formula
    elif value is not None:
        cell.value = value
    if fill:   cell.fill      = fill
    if font:   cell.font      = font
    if align:  cell.alignment = align
    if border: cell.border    = border
    if fmt:    cell.number_format = fmt
    return cell


# ── MAIN FUNCTION ────────────────────────────────────
def generate_pallet_excel(rfq, results, edited_data=None):
    """
    Generate pallet costing Excel.
    Handles both plywood deck (SQM) and pinewood deck (CFT).
    """
    wb = Workbook()
    wb.remove(wb.active)

    for idx, r in enumerate(results):
        sheet_name = f"Box {idx+1}"[:31]
        ws = wb.create_sheet(title=sheet_name)
        _build_pallet_sheet(
            ws, rfq,
            r['box'], r['design'],
            r['bom'],  r['totals'],
            edited_data=edited_data
        )

    output_dir = settings.OUTPUT_DIR
    os.makedirs(output_dir, exist_ok=True)
    filename = (
        f"{rfq.rfq_number.replace('/', '-')}"
        f"_pallet_costing.xlsx"
    )
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    return filepath


def _build_pallet_sheet(ws, rfq, box, design,
                         bom, totals, edited_data=None):

    # ── COLUMN WIDTHS ────────────────────────────────
    widths = {
        'A': 7,  'B': 24, 'C': 20,
        'D': 12, 'E': 12, 'F': 12,
        'G': 10, 'H': 13, 'I': 8,
        'J': 12, 'K': 14,
        'L': 3,  'M': 14, 'N': 10,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # ── ROW 1: COMPANY ───────────────────────────────
    ws.merge_cells('A1:K1')
    sc(ws, 1, 1,
       value='PRONK MULTISERVICE INDIA PVT. LTD.',
       fill=hfill(RED),
       font=Font(bold=True, size=14,
                 color=WHITE, name='Calibri'),
       align=halign('center', 'center'))
    ws.row_dimensions[1].height = 30

    # ── ROW 2: BRANCH / DATE ─────────────────────────
    ws.row_dimensions[2].height = 18
    sc(ws, 2, 1, value='BRANCH:',
       font=hfont(bold=True, size=10))
    sc(ws, 2, 2, value='Bengaluru',
       font=hfont(bold=False, size=10))
    sc(ws, 2, 8, value='Date:',
       font=hfont(bold=True, size=10))
    from datetime import date
    sc(ws, 2, 9,
       value=date.today().strftime('%d-%m-%Y'),
       font=hfont(bold=False, size=10))
    sc(ws, 2, 13, value='W',
       font=hfont(bold=True, size=10),
       align=halign())
    sc(ws, 2, 14, value='H',
       font=hfont(bold=True, size=10),
       align=halign())

    # ── ROW 3: CUSTOMER ──────────────────────────────
    ws.row_dimensions[3].height = 18
    sc(ws, 3, 1, value='CUSTOMER :',
       font=hfont(bold=True, size=10))
    client = rfq.client.company_name if rfq.client else ''
    sc(ws, 3, 2, value=client,
       font=hfont(bold=False, size=10))

    # Sidebar: Pallet Deck
    # For plywood deck → show OD_W as width
    # For pinewood deck → show deck_w (plank width)
    is_plywood = design.get('deck_uom', 'CFT') == 'SQM'
    deck_w_display = (
        design.get('box_od_w', 0) if is_plywood
        else design.get('deck_w', 150)
    )
    deck_h_display = design.get('deck_h', 8)

    sc(ws, 3, 12, value='Pallet Deck',
       font=hfont(bold=True, size=9),
       align=halign('right'))
    sc(ws, 3, 13, value=deck_w_display,
       font=hfont(bold=False, size=10),
       align=halign())
    sc(ws, 3, 14, value=deck_h_display,
       font=hfont(bold=False, size=10),
       align=halign())

    # ── ROW 4: PRODUCT + BOX ID ──────────────────────
    # KEY: E4=ID_L, F4=ID_W, G4=ID_H
    ws.row_dimensions[4].height = 20
    sc(ws, 4, 1, value='PRODUCT NAME:',
       font=hfont(bold=True, size=10))
    sc(ws, 4, 2, value=box.product_name,
       font=hfont(bold=False, size=10))
    sc(ws, 4, 4, value='BOX ID',
       fill=hfill(DARK_BLUE),
       font=Font(bold=True, size=10,
                 color=WHITE, name='Calibri'),
       align=halign())
    # E4=ID_L, F4=ID_W, G4=ID_H
    sc(ws, 4, 5, value=box.box_id_length,
       font=hfont(bold=True, size=10),
       align=halign(), fill=hfill(LIGHT_BLUE))
    sc(ws, 4, 6, value=box.box_id_width,
       font=hfont(bold=True, size=10),
       align=halign(), fill=hfill(LIGHT_BLUE))
    sc(ws, 4, 7, value=box.box_id_height,
       font=hfont(bold=True, size=10),
       align=halign(), fill=hfill(LIGHT_BLUE))
    sc(ws, 4, 8, value='Volume(CBM)',
       font=hfont(bold=True, size=9),
       align=halign('right'))
    sc(ws, 4, 10, value=design['volumes']['cbm'],
       font=hfont(bold=True, size=10),
       align=halign(), fmt='0.00')

    # Sidebar: Pallet w Runner
    sc(ws, 4, 12, value='Pallet w Runner',
       font=hfont(bold=True, size=9),
       align=halign('right'))
    sc(ws, 4, 13, value=design.get('runner_w', 100),
       font=hfont(bold=False, size=10),
       align=halign())
    sc(ws, 4, 14, value=design.get('runner_h', 100),
       font=hfont(bold=False, size=10),
       align=halign())

    # ── ROW 5: PALLET OD ─────────────────────────────
    # KEY: E5=OD_L, F5=OD_W, G5=OD_H ← BOM uses these!
    ws.row_dimensions[5].height = 20
    sc(ws, 5, 4, value='PALLET OD',
       fill=hfill(RED),
       font=Font(bold=True, size=10,
                 color=WHITE, name='Calibri'),
       align=halign())
    # E5=OD_L, F5=OD_W, G5=OD_H
    sc(ws, 5, 5, value=design['box_od_l'],
       font=Font(bold=True, size=10,
                 color='C00000', name='Calibri'),
       align=halign(), fill=hfill(YELLOW))
    sc(ws, 5, 6, value=design['box_od_w'],
       font=Font(bold=True, size=10,
                 color='C00000', name='Calibri'),
       align=halign(), fill=hfill(YELLOW))
    sc(ws, 5, 7, value=design['box_od_h'],
       font=Font(bold=True, size=10,
                 color='C00000', name='Calibri'),
       align=halign(), fill=hfill(YELLOW))
    sc(ws, 5, 8, value='Total Area(SQM)',
       font=hfont(bold=True, size=9),
       align=halign('right'))
    sc(ws, 5, 10,
       value=design['volumes']['total_area_sqm'],
       font=hfont(bold=True, size=10),
       align=halign(), fmt='0.00')

    # Sidebar: Walls (pallet has none)
    sc(ws, 5, 12, value='Walls',
       font=hfont(bold=True, size=9),
       align=halign('right'))
    sc(ws, 5, 13, value=0,
       font=hfont(bold=False, size=10),
       align=halign())

    # ── ROW 6: BASE AREA ─────────────────────────────
    ws.row_dimensions[6].height = 18
    sc(ws, 6, 8, value='Base Area(SQM)',
       font=hfont(bold=True, size=9),
       align=halign('right'))
    sc(ws, 6, 10,
       value=design['volumes']['base_area_sqm'],
       font=hfont(bold=True, size=10),
       align=halign(), fmt='0.00')
    sc(ws, 6, 12, value='Beading',
       font=hfont(bold=True, size=9),
       align=halign('right'))
    sc(ws, 6, 13, value=0,
       font=hfont(bold=False, size=10),
       align=halign())

    # ── ROW 7: QUANTITY / WEIGHT ─────────────────────
    ws.row_dimensions[7].height = 18
    sc(ws, 7, 6, value='Quantity',
       font=hfont(bold=True, size=10))
    sc(ws, 7, 7, value=box.quantity,
       font=hfont(bold=False, size=10),
       align=halign())
    sc(ws, 7, 8, value='Weight',
       font=hfont(bold=True, size=10))
    sc(ws, 7, 9,
       value=f"{box.weight_kg} Kgs",
       font=hfont(bold=False, size=10))

    # ── ROW 8 ────────────────────────────────────────
    ws.row_dimensions[8].height = 18
    sc(ws, 8, 1, value='DRAWING AVL(Y/N):',
       font=hfont(bold=True, size=9))
    sc(ws, 8, 2, value='N',
       font=hfont(bold=False, size=10))

    # ── ROW 9: COSTING BY ────────────────────────────
    ws.row_dimensions[9].height = 18
    sc(ws, 9, 1, value='COSTING BY:',
       font=hfont(bold=True, size=9))
    eng = (rfq.engineer.full_name or rfq.engineer.username
           if rfq.engineer else '')
    sc(ws, 9, 2, value=eng,
       font=hfont(bold=False, size=10))

    # ── ROW 10: APPROVED BY ──────────────────────────
    ws.row_dimensions[10].height = 18
    sc(ws, 10, 1, value='APPROVED BY:',
       font=hfont(bold=True, size=9))

    # ── ROW 11: COLUMN HEADERS ───────────────────────
    ws.row_dimensions[11].height = 30
    ws.merge_cells('D11:F11')
    headers = [
        (1,  'SR. NO.'),
        (2,  'DESCRIPTION'),
        (3,  'MATERIAL'),
        (4,  'DIMENSIONS in mm'),
        (7,  'QTY IN\nNOS.'),
        (8,  'TOTAL QTY'),
        (9,  'UOM'),
        (10, 'RATE/UOM'),
        (11, 'TOTAL'),
    ]
    for col_num, text in headers:
        sc(ws, 11, col_num,
           value=text,
           fill=hfill(RED),
           font=Font(bold=True, size=9,
                     color=WHITE, name='Calibri'),
           align=halign('center', 'center', wrap=True),
           border=hborder())

    # ── ROW 12: L/W/H SUBHEADERS ─────────────────────
    ws.row_dimensions[12].height = 16
    for col_num, text in [(4,'L'),(5,'W'),(6,'H/T')]:
        sc(ws, 12, col_num,
           value=text,
           fill=hfill(RED),
           font=Font(bold=True, size=9,
                     color=WHITE, name='Calibri'),
           align=halign(),
           border=hborder())

    # ── ROW 13: PALLET LABEL ─────────────────────────
    ws.row_dimensions[13].height = 16
    ws.merge_cells('A13:K13')
    sc(ws, 13, 1,
       value='Pallet / Skid',
       fill=hfill(TEAL),
       font=Font(bold=True, size=10,
                 color=WHITE, name='Calibri'),
       align=halign())

    # ── KEY CELL REFERENCES ──────────────────────────
    # E5=OD_L, F5=OD_W, G5=OD_H, G4=ID_H
    OD_L = 'E5'
    OD_W = 'F5'
    OD_H = 'G5'
    ID_H = 'G4'
    CBM  = 'J4'

    rule = design.get('rule')

    # ── BOM ROWS (start row 14) ──────────────────────
    BOM_START = 14

    for i, item in enumerate(bom):
        r   = BOM_START + i
        ws.row_dimensions[r].height = 16
        desc = item['description']
        uom  = item['uom'].upper()

        # Use edited values if provided
        if edited_data and i < len(edited_data):
            ed      = edited_data[i]
            qty_nos = ed[6] if ed[6] else item['qty_nos']
            rate    = ed[9] if ed[9] else item['rate_per_uom']
            l_val   = ed[3] if ed[3] else item['length_mm']
            w_val   = ed[4] if ed[4] else item['width_mm']
            h_val   = ed[5] if ed[5] else item['thickness_mm']
        else:
            qty_nos = item['qty_nos']
            rate    = item['rate_per_uom']
            l_val   = item['length_mm']
            w_val   = item['width_mm']
            h_val   = item['thickness_mm']

        # ── SR ───────────────────────────────────────
        sc(ws, r, 1, value=item['sr_no'],
           align=halign(), border=hborder(),
           font=hfont(bold=False, size=9))

        # ── DESCRIPTION ──────────────────────────────
        sc(ws, r, 2, value=desc,
           align=Alignment(horizontal='left',
                           vertical='center'),
           border=hborder(),
           font=hfont(bold=False, size=9))

        # ── MATERIAL ─────────────────────────────────
        sc(ws, r, 3, value=item['material'],
           align=Alignment(horizontal='left',
                           vertical='center'),
           border=hborder(),
           font=hfont(bold=False, size=9))

        # ── L COLUMN ─────────────────────────────────
        # Pallet L column formula map
        if 'Deck' in desc:
            if is_plywood:
                # Plywood deck → L = OD_L
                l_formula = f'={OD_L}'
            else:
                # Pinewood deck → L = deck_length
                # If od_l > 2500 → deck runs along width
                if design.get('box_od_l', 0) > 2500:
                    l_formula = f'={OD_W}'
                else:
                    l_formula = f'={OD_L}'
            sc(ws, r, 4, formula=l_formula,
               align=halign(), border=hborder(),
               fill=hfill(GREEN),
               font=hfont(bold=False, size=9),
               fmt='0')

        elif 'Width Runner' in desc or 'Bottom Chocking' in desc:
            sc(ws, r, 4, formula=f'={OD_W}',
               align=halign(), border=hborder(),
               fill=hfill(GREEN),
               font=hfont(bold=False, size=9),
               fmt='0')

        elif 'Length Runner' in desc:
            sc(ws, r, 4, formula=f'={OD_L}',
               align=halign(), border=hborder(),
               fill=hfill(GREEN),
               font=hfont(bold=False, size=9),
               fmt='0')

        else:
            # Hardware / consumable — no L dimension
            sc(ws, r, 4, value='',
               border=hborder())

        # ── W COLUMN ─────────────────────────────────
        if 'Deck' in desc and is_plywood:
            # Plywood deck → W = OD_W (full pallet width)
            sc(ws, r, 5, formula=f'={OD_W}',
               align=halign(), border=hborder(),
               fill=hfill(GREEN),
               font=hfont(bold=False, size=9),
               fmt='0')
        else:
            sc(ws, r, 5,
               value=w_val if w_val else '',
               align=halign(), border=hborder(),
               font=hfont(bold=False, size=9),
               fmt='0')

        # ── H/T COLUMN ───────────────────────────────
        sc(ws, r, 6,
           value=h_val if h_val else '',
           align=halign(), border=hborder(),
           font=hfont(bold=False, size=9),
           fmt='0')

        # ── QTY NOS (ROUNDUP formula!) ───────────────
        qty_formula = _get_pallet_qty_formula(
            desc, rule, OD_L, OD_W, OD_H, ID_H, CBM,
            is_plywood
        )

        if qty_formula:
            sc(ws, r, 7, formula=qty_formula,
               align=halign(), border=hborder(),
               fill=hfill(GREEN),
               font=hfont(bold=False, size=9),
               fmt='0')
        else:
            sc(ws, r, 7, value=qty_nos,
               align=halign(), border=hborder(),
               font=hfont(bold=False, size=9),
               fmt='0.00')

        # ── TOTAL QTY ────────────────────────────────
        if uom == 'SQM' and 'Deck' in desc and is_plywood:
            # Plywood deck SQM = OD_L/1000 × OD_W/1000 × 1.05
            tq_formula = (
                f'=({OD_L}/1000)*'
                f'({OD_W}/1000)*1*1.05'
            )
            sc(ws, r, 8, formula=tq_formula,
               align=halign('right'), border=hborder(),
               fill=hfill(GREEN),
               font=hfont(bold=False, size=9),
               fmt='0.00')

        elif uom == 'CFT' and l_val:
            # Pinewood CFT formula
            tq_formula = (
                f'=((D{r}/1000)*(E{r}/1000)'
                f'*(F{r}/1000))'
                f'*G{r}*(3.28^3)*1.05'
            )
            sc(ws, r, 8, formula=tq_formula,
               align=halign('right'), border=hborder(),
               fill=hfill(GREEN),
               font=hfont(bold=False, size=9),
               fmt='0.00')

        elif uom == 'SQM' and l_val and 'Deck' not in desc:
            # Other SQM items
            tq_formula = (
                f'=(D{r}/1000)*(E{r}/1000)'
                f'*G{r}*1.05'
            )
            sc(ws, r, 8, formula=tq_formula,
               align=halign('right'), border=hborder(),
               fill=hfill(GREEN),
               font=hfont(bold=False, size=9),
               fmt='0.00')

        else:
            sc(ws, r, 8,
               value=item['total_qty'],
               align=halign('right'), border=hborder(),
               font=hfont(bold=False, size=9),
               fmt='0.00')

        # ── UOM ──────────────────────────────────────
        sc(ws, r, 9, value=uom,
           align=halign(), border=hborder(),
           font=hfont(bold=False, size=9))

        # ── RATE ─────────────────────────────────────
        sc(ws, r, 10, value=rate,
           align=halign('right'), border=hborder(),
           font=hfont(bold=False, size=9),
           fmt='#,##0')

        # ── TOTAL formula ────────────────────────────
        if uom == 'MTR' or 'Silica' in desc:
            total_formula = f'=G{r}*H{r}*J{r}'
        else:
            total_formula = f'=H{r}*J{r}'

        sc(ws, r, 11, formula=total_formula,
           align=halign('right'), border=hborder(),
           fill=hfill(GREEN),
           font=Font(bold=True, size=9,
                     color='1B5E20', name='Calibri'),
           fmt='#,##0')

    # ── TOTALS SECTION ───────────────────────────────
    last_row = BOM_START + len(bom) - 1
    K_range  = f'K{BOM_START}:K{last_row}'
    t        = last_row + 2

    def tot_row_add(label, formula, bg, row):
        ws.merge_cells(f'A{row}:J{row}')
        sc(ws, row, 1, value=label,
           fill=hfill(bg),
           font=Font(
               bold=True, size=10,
               color=WHITE if bg in [RED, DARK_BLUE, TEAL]
               else '000000',
               name='Calibri'),
           align=halign('right'))
        sc(ws, row, 11, formula=formula,
           fill=hfill(bg),
           font=Font(
               bold=True, size=10,
               color=WHITE if bg in [RED, DARK_BLUE, TEAL]
               else '000000',
               name='Calibri'),
           align=halign('right'), fmt='#,##0')
        ws.row_dimensions[row].height = 20

    tot_row_add('TOTAL COST',
                f'=SUM({K_range})',
                DARK_BLUE, t)
    t += 1
    tot_row_add(
        f'Overhead @ {totals["overhead_pct"]}%',
        f'=K{t-1}*{totals["overhead_pct"]}/100',
        ORANGE, t)
    t += 1
    tot_row_add('After Overhead',
                f'=K{t-2}+K{t-1}',
                GRAY, t)
    t += 1
    tot_row_add(
        f'Margin @ {totals["margin_pct"]}%',
        f'=K{t-1}*{totals["margin_pct"]}/100',
        ORANGE, t)
    t += 1
    tot_row_add(
        'Before GST (CEILING to 100)',
        f'=CEILING(K{t-2}+K{t-1},100)',
        GRAY, t)
    t += 1
    tot_row_add(
        f'GST @ {totals["gst_pct"]}%',
        f'=K{t-1}*{totals["gst_pct"]}/100',
        ORANGE, t)
    t += 1
    tot_row_add('GRAND TOTAL',
                f'=K{t-2}+K{t-1}',
                RED, t)

    # ── FREEZE + PRINT ───────────────────────────────
    ws.freeze_panes = 'A14'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1


def _get_pallet_qty_formula(desc, rule, OD_L, OD_W,
                             OD_H, ID_H, CBM, is_plywood):
    """ROUNDUP formulas for pallet BOM"""
    def p(attr, default):
        return getattr(rule, attr) if rule else default

    if 'Deck' in desc:
        if is_plywood:
            return '=1'  # always 1 sheet for plywood
        else:
            # Pinewood deck qty by pitch
            pitch = p('deck_pitch', 140)
            return f'=ROUNDUP({OD_L}/{pitch},0)'

    elif 'Length Runner' in desc:
        pitch = p('l_runner_pitch', 500)
        return f'=ROUNDUP({OD_W}/{pitch},0)'

    elif 'Width Runner' in desc:
        pitch = p('w_runner_pitch', 600)
        return f'=ROUNDUP({OD_L}/{pitch},0)'

    elif 'Bottom Chocking' in desc:
        pitch = p('bot_chock_pitch', 500)
        return f'=ROUNDUP({OD_W}/{pitch},0)'

    elif 'Top Chocking' in desc:
        pitch = p('top_chock_pitch', 500)
        return f'=ROUNDUP({OD_L}/{pitch},0)'

    elif 'Lashing' in desc:
        pitch = p('lashing_pitch', 600)
        return f'=ROUNDUP({OD_L}/{pitch},0)'

    elif 'Silica' in desc:
        return f'=ROUNDUP({CBM}*500/50,0)'

    elif 'N/B' in desc or 'Nut' in desc:
        return (f'=ROUNDUP(({OD_L}/250)*4+'
                f'({OD_W}/250)*4+'
                f'({OD_H}/250)*4,0)')

    return None